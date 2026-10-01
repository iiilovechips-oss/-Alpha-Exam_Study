"""Turn a source file into what NotebookLM can take: Excel becomes Markdown, big files are split."""

import io
import math
import shutil
from pathlib import Path

PART_MB = 8          # larger PDFs and slide decks are split into parts no bigger than this
MAX_ROWS = 60        # rows of each sheet shown in the Markdown version of a workbook
MAX_COLS = 30
MAX_FORMULAS = 60    # formulas listed per sheet
FORMULA_SCAN_ROWS = 500


def render(src: Path, out_dir: Path, name: str, cache: Path | None = None) -> list[str]:
    """Write `src` into `out_dir` under `name` (or several derived names) and return the names.

    Conversions are slow for big decks, so results are kept in `cache` (one folder per source
    version) and reused until the source file changes.
    """
    suffix = src.suffix.lower()
    big = src.stat().st_size > PART_MB * 1_000_000 and suffix in (".pdf", ".pptx")
    if suffix != ".xlsx" and not big:
        shutil.copy2(src, out_dir / name)
        return [name]

    stem = Path(name).stem
    if cache is None or not cache.exists():
        made = {}
        if suffix == ".xlsx":
            made[f"{stem} (Excel).md"] = xlsx_to_markdown(src).encode()
        else:
            parts = split(src, pdf_part if suffix == ".pdf" else pptx_part)
            for i, data in enumerate(parts, 1):
                made[f"{stem} (part {i} of {len(parts)}){suffix}"] = data
        target = cache or out_dir
        target.mkdir(parents=True, exist_ok=True)
        for made_name, data in made.items():
            (target / made_name).write_bytes(data)
        if cache is None:
            if not made:
                shutil.copy2(src, out_dir / name)
            return sorted(made) or [name]

    names = sorted(p.name for p in cache.iterdir())
    if not names:  # could not be split: use the file whole
        shutil.copy2(src, out_dir / name)
        return [name]
    for made_name in names:
        shutil.copy2(cache / made_name, out_dir / made_name)
    return names


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        value = round(value, 4)
    return str(value).replace("|", "\\|").replace("\n", " ")


def xlsx_to_markdown(src: Path) -> str:
    """Values as tables plus the formulas behind them, so the logic of the workbook survives."""
    from openpyxl import load_workbook

    values = load_workbook(src, data_only=True, read_only=True)
    formulas = load_workbook(src, read_only=True)
    lines = [f"# {src.name}", "", "Converted from Excel: cell values shown as tables, formulas listed under each sheet.", ""]
    for sheet in values.worksheets:
        rows = [list(r[:MAX_COLS]) for r in sheet.iter_rows(max_row=MAX_ROWS, values_only=True)]
        rows = [r for r in rows if any(c is not None for c in r)]
        lines += [f"## Sheet: {sheet.title}", ""]
        if not rows:
            lines += ["(empty)", ""]
            continue
        width = max(len(r) for r in rows)
        table = [[cell_text(c) for c in r] + [""] * (width - len(r)) for r in rows]
        lines.append("| " + " | ".join(table[0]) + " |")
        lines.append("|" + " --- |" * width)
        lines += ["| " + " | ".join(r) + " |" for r in table[1:]]
        total = sheet.max_row or len(rows)
        if total > MAX_ROWS:
            lines += ["", f"(first {MAX_ROWS} of {total} rows shown)"]
        found = []
        for row in formulas[sheet.title].iter_rows(max_row=FORMULA_SCAN_ROWS):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=") and len(found) < MAX_FORMULAS:
                    found.append(f"- {cell.coordinate}: `{cell.value}`")
        if found:
            lines += ["", "Formulas:"] + found
        lines.append("")
    return "\n".join(lines)


def split(src: Path, make_part) -> list[bytes]:
    """Split into the fewest equal page/slide ranges whose parts all fit under PART_MB."""
    limit = PART_MB * 1_000_000
    count = make_part(src, None)
    pieces = max(2, math.ceil(src.stat().st_size / limit))
    while pieces <= count:
        size = math.ceil(count / pieces)
        parts = [make_part(src, range(start, min(start + size, count))) for start in range(0, count, size)]
        if all(len(p) <= limit for p in parts):
            return parts
        pieces += 1
    return []  # even single pages are too big; the caller copies the file whole


def pdf_part(src: Path, pages: range | None):
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(src)
    if pages is None:
        return len(reader.pages)
    writer = PdfWriter()
    for i in pages:
        writer.add_page(reader.pages[i])
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def pptx_part(src: Path, slides: range | None):
    from pptx import Presentation

    deck = Presentation(src)
    ids = deck.slides._sldIdLst
    if slides is None:
        return len(ids)
    for index, slide_id in enumerate(list(ids)):
        if index not in slides:
            deck.part.drop_rel(slide_id.rId)  # unreferenced slides and their media are not saved
            ids.remove(slide_id)
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()
